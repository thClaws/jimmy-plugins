#!/usr/bin/env python3
"""slides2preview.py — one-shot render + preview + font-proof pipeline.

chNN.yaml ──render──▶ native .pptx ──soffice──▶ .pdf ──pdftoppm──▶ slide PNGs
                                                          │
                              contact sheet (QC) ◀────────┤
                              image-pptx (font-proof, notes preserved) ◀──┘

WHY: a .pptx stores only font *names*. Every viewer (PowerPoint on Mac, the Cowork
file preview, LibreOffice, Keynote) substitutes missing fonts differently, so the
same deck looks different everywhere. This script renders ONCE with the fonts
installed on this machine, verifies those fonts actually resolve (no silent
fallback), and bakes each slide into an image inside a pptx — identical on every
machine, no font install needed. Speaker notes are copied from the YAML so the
TTS/video pipeline still works.

Usage:
  python3 slides2preview.py media/slides/ch01.yaml \
      [--book book.yaml] [--theme theme.json] [--media-dir DIR ...] \
      [--out build/slides] [--dpi 150] [--strict-fonts] \
      [--skip-image-pptx] [--skip-contact] [--skip-native]

Outputs for input chNN.yaml (in --out, default build/slides):
  chNN.pptx                      native deck (editable; correct only where fonts exist)
  chNN-slides.pdf                exact-render PDF (fonts embedded — safe to share)
  preview/chNN/slide-*.png       per-slide preview images
  preview/chNN-contact.png       contact sheet for quick visual QC
  chNN-image.pptx                font-proof deck: each slide = baked image + notes

Requires: python-pptx, pyyaml, Pillow, fontconfig (fc-match),
          soffice (libreoffice), pdftoppm (poppler-utils).
"""
import argparse, glob, math, os, shutil, subprocess, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import slides2pptx  # noqa: E402  (reuse renderer + theme loader)

try:
    import yaml
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
except ImportError as e:
    sys.stderr.write(f"missing dependency: {e}\n"
                     "pip install python-pptx pyyaml Pillow --break-system-packages\n")
    sys.exit(2)

def log(msg):
    sys.stderr.write(f"[slides2preview] {msg}\n")

def need_tool(name, hint):
    if shutil.which(name) is None:
        log(f"ERROR: `{name}` not found — {hint}")
        sys.exit(2)

def check_fonts(theme, strict):
    """Guardrail: verify every theme font actually resolves via fontconfig.
    A missing font does NOT fail the render — it silently substitutes, which is
    exactly how broken decks get shipped. Catch it here, loudly."""
    if shutil.which("fc-match") is None:
        log("WARNING: fc-match not available — cannot verify fonts, output may substitute")
        return
    bad = []
    for role, fam in (theme.get("fonts") or {}).items():
        got = subprocess.run(["fc-match", "-f", "%{family}", fam],
                             capture_output=True, text=True).stdout
        families = [f.strip().lower() for f in got.split(",")]
        if fam.strip().lower() not in families:
            bad.append((role, fam, got))
    if bad:
        for role, fam, got in bad:
            log(f"FONT MISMATCH: {role} wants '{fam}' but system resolves to '{got}'")
        if strict:
            log("aborting (--strict-fonts). Install the fonts, run fc-cache -f, retry.")
            sys.exit(3)
        log("WARNING: continuing with substituted fonts — preview will NOT match design")
    else:
        log("fonts OK: " + ", ".join(f"{r}={f}" for r, f in (theme.get("fonts") or {}).items()))

def pptx_to_pdf(pptx_path, outdir):
    profile = tempfile.mkdtemp(prefix="lo_profile_")
    cmd = ["soffice", "--headless", "--norestore",
           f"-env:UserInstallation=file://{profile}",
           "--convert-to", "pdf", "--outdir", outdir, pptx_path]
    for attempt in (1, 2):  # soffice is occasionally flaky on first launch
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
        pdf = os.path.join(outdir, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
        if r.returncode == 0 and os.path.exists(pdf):
            return pdf
        log(f"soffice attempt {attempt} failed: {r.stderr.strip()[:200]}")
    sys.exit(4)

def pdf_to_pngs(pdf, png_dir, dpi):
    os.makedirs(png_dir, exist_ok=True)
    for f in glob.glob(os.path.join(png_dir, "slide-*.png")):
        os.remove(f)
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), pdf,
                    os.path.join(png_dir, "slide")], check=True, timeout=600)
    return sorted(glob.glob(os.path.join(png_dir, "slide-*.png")))

def contact_sheet(pngs, out_path, cols=4, thumb_w=320, pad=8, bg=(10, 21, 18)):
    ims = []
    for p in pngs:
        im = Image.open(p)
        h = round(im.height * thumb_w / im.width)
        ims.append(im.resize((thumb_w, h), Image.LANCZOS))
    th = ims[0].height
    rows = math.ceil(len(ims) / cols)
    sheet = Image.new("RGB", (cols * thumb_w + (cols + 1) * pad,
                              rows * th + (rows + 1) * pad), bg)
    for i, im in enumerate(ims):
        r, c = divmod(i, cols)
        sheet.paste(im, (pad + c * (thumb_w + pad), pad + r * (th + pad)))
    sheet.save(out_path)

def build_image_pptx(pngs, slides_meta, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for img, s in zip(pngs, slides_meta):
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        sl.shapes.add_picture(img, 0, 0, prs.slide_width, prs.slide_height)
        if s.get("notes"):
            sl.notes_slide.notes_text_frame.text = str(s["notes"])
    prs.save(out_path)

def main(argv=None):
    ap = argparse.ArgumentParser(description="chNN.yaml -> pptx + pdf + preview PNGs + image-pptx")
    ap.add_argument("input", help="path to compiled chNN.yaml (never .slides.md)")
    ap.add_argument("--book", help="book.yaml (style.pptx_theme + media dirs)")
    ap.add_argument("--theme", help="theme.json deep-merged over default")
    ap.add_argument("--media-dir", action="append", default=[])
    ap.add_argument("--out", default="build/slides")
    ap.add_argument("--dpi", type=int, default=150)
    ap.add_argument("--strict-fonts", action="store_true",
                    help="fail (exit 3) if any theme font is not actually installed")
    ap.add_argument("--skip-image-pptx", action="store_true")
    ap.add_argument("--skip-contact", action="store_true")
    ap.add_argument("--skip-native", action="store_true",
                    help="don't keep the native pptx in --out")
    args = ap.parse_args(argv)

    need_tool("soffice", "apt-get install libreoffice  (or libreoffice-impress)")
    need_tool("pdftoppm", "apt-get install poppler-utils")

    data = yaml.safe_load(open(args.input, encoding="utf-8"))
    if not data or "slides" not in data:
        log("input has no 'slides'"); return 1
    slides_meta = data["slides"]
    stem = os.path.splitext(os.path.basename(args.input))[0]
    os.makedirs(args.out, exist_ok=True)

    # 0) font guardrail
    theme = slides2pptx.load_theme(args.book, args.theme)
    check_fonts(theme, args.strict_fonts)

    # 1) native pptx (reuse the standalone renderer)
    native = os.path.join(args.out, f"{stem}.pptx")
    rc = slides2pptx.main([args.input, "-o", native]
                          + (["--book", args.book] if args.book else [])
                          + (["--theme", args.theme] if args.theme else [])
                          + sum([["--media-dir", d] for d in args.media_dir], []))
    if rc not in (0, None):
        return rc

    # 2) exact-render PDF
    with tempfile.TemporaryDirectory(prefix="s2p_") as tmp:
        pdf_tmp = pptx_to_pdf(native, tmp)
        pdf_out = os.path.join(args.out, f"{stem}-slides.pdf")
        shutil.copy(pdf_tmp, pdf_out)
    log(f"pdf   -> {pdf_out}")

    # 3) per-slide preview PNGs
    png_dir = os.path.join(args.out, "preview", stem)
    pngs = pdf_to_pngs(pdf_out, png_dir, args.dpi)
    if len(pngs) != len(slides_meta):
        log(f"ERROR: {len(pngs)} PNGs != {len(slides_meta)} slides"); return 5
    log(f"pngs  -> {png_dir} ({len(pngs)} slides @{args.dpi}dpi)")

    # 4) contact sheet for QC
    if not args.skip_contact:
        sheet = os.path.join(args.out, "preview", f"{stem}-contact.png")
        contact_sheet(pngs, sheet)
        log(f"sheet -> {sheet}")

    # 5) font-proof image-pptx with speaker notes
    if not args.skip_image_pptx:
        img_pptx = os.path.join(args.out, f"{stem}-image.pptx")
        build_image_pptx(pngs, slides_meta, img_pptx)
        n_notes = sum(1 for s in slides_meta if s.get("notes"))
        log(f"image-pptx -> {img_pptx} ({len(pngs)} slides, {n_notes} with notes)")

    if args.skip_native:
        os.remove(native)
    log("done")
    return 0

if __name__ == "__main__":
    sys.exit(main())
